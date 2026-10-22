import io
import json
from datetime import date, datetime, timedelta
from urllib.parse import unquote

import pytest
from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.imports_v35 import COOKIE, ContactImportBatch, router
from app.models import Lead
from app.services import contact_import as service


@pytest.fixture
def setup(tmp_path):
    engine = create_engine('sqlite:///' + str(tmp_path / 'test.db'), connect_args={'check_same_thread': False})
    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine)
    app = FastAPI()
    app.mount("/static", StaticFiles(directory="app/static"), name="static")
    app.include_router(router)
    def db():
        with sessions() as session:
            yield session
    app.dependency_overrides[get_db] = db
    with TestClient(app) as client:
        client.get('/imports')
        yield client, sessions
    engine.dispose()


def stage(client, content, filename='contacts.csv', worksheet=''):
    response = client.post('/imports/source', data={'csrf': client.cookies[COOKIE], 'source_type': 'file', 'worksheet': worksheet}, files={'file': (filename, content)}, follow_redirects=False)
    assert response.status_code == 303, response.text
    return response.headers['location']


def preview(client, url, fields, mode='skip'):
    return client.post(url + '/preview', data={'csrf': client.cookies[COOKIE], 'mode': mode, **{f'column_{i}': field for i, field in enumerate(fields)}})


def commit(client, url):
    return client.post(url + '/commit', data={'csrf': client.cookies[COOKIE], 'confirm': 'yes'})


def test_csv_mapping_validation_and_idempotent_commit(setup):
    client, sessions = setup
    url = stage(client, b'First Name,Last Name,Email,DOB,ZIP,PDP\nAnn,Example,ANN@example.test,04/15/1961,501,Yes\nAnn,Example,ann@example.test,04/15/1961,501,Yes\nBad,Date,bad@example.test,02/30/1960,97467,no\n')
    assert 'Match your columns' in client.get(url).text
    result = preview(client, url, ['first_name', 'last_name', 'email', 'date_of_birth', 'zip_code', 'interested_pdp'])
    assert result.status_code == 200
    assert 'First Name' in result.text
    with sessions() as db:
        assert db.scalars(select(Lead)).all() == []
        stored = json.loads(db.get(ContactImportBatch, url.rsplit('/', 1)[-1]).payload)['preview']
        assert stored['counts'] == {'create': 1, 'update': 0, 'skip': 1, 'error': 1}
    assert commit(client, url).status_code == 200
    assert commit(client, url).status_code == 200
    with sessions() as db:
        leads = db.scalars(select(Lead)).all()
        assert len(leads) == 1
        assert leads[0].email == 'ann@example.test'
        assert leads[0].date_of_birth == date(1961, 4, 15)
        assert leads[0].zip_code == '00501'
        assert leads[0].interested_pdp is True
        batch = db.get(ContactImportBatch, url.rsplit('/', 1)[-1])
        assert 'table' not in json.loads(batch.payload)
    assert 'Source row,Result,Details' in client.get(url + '/results.csv').text


def test_update_blanks_false_and_followup(setup):
    client, sessions = setup
    with sessions.begin() as db:
        db.add(Lead(first_name='Ann', last_name='Example', email='ann@example.test', phone='555-555-0100', notes='Keep this', interested_pdp=True))
    url = stage(client, b'First,Last,Email,Notes,PDP,Next Follow-up\nAnn,Example,ann@example.test,,false,2026-11-03\n')
    preview(client, url, ['first_name', 'last_name', 'email', 'notes', 'interested_pdp', 'next_follow_up'], 'update')
    assert commit(client, url).status_code == 200
    with sessions() as db:
        lead = db.scalar(select(Lead))
        assert lead.notes == 'Keep this'
        assert lead.phone == '555-555-0100'
        assert lead.interested_pdp is False
        assert lead.next_follow_up == date(2026, 11, 3)


def test_skip_existing_default(setup):
    client, sessions = setup
    with sessions.begin() as db:
        db.add(Lead(first_name='Ann', last_name='Example', email='ann@example.test', notes='Original'))
    url = stage(client, b'First,Last,Email,Notes\nAnn,Example,ann@example.test,Changed\n')
    preview(client, url, ['first_name', 'last_name', 'email', 'notes'])
    commit(client, url)
    with sessions() as db:
        assert db.scalar(select(Lead)).notes == 'Original'


def test_stale_preview_requires_reconfirmation(setup):
    client, sessions = setup
    url = stage(client, b'First,Last,Email\nAnn,Example,ann@example.test\n')
    preview(client, url, ['first_name', 'last_name', 'email'])
    with sessions.begin() as db:
        db.add(Lead(first_name='Ann', last_name='Example', email='ann@example.test'))
    response = commit(client, url)
    assert 'Contacts changed since your preview' in response.text
    with sessions() as db:
        assert len(db.scalars(select(Lead)).all()) == 1
        assert db.get(ContactImportBatch, url.rsplit('/', 1)[-1]).state == 'preview'


def test_duplicates_conflicts_and_dnc(setup):
    client, sessions = setup
    with sessions.begin() as db:
        db.add_all([Lead(first_name='Ann', last_name='Example', email='ann@example.test', phone='555-555-0100', status='Do Not Call'), Lead(first_name='Bob', last_name='Example', email='bob@example.test', phone='555-555-0200')])
    url = stage(client, b'First,Last,Email,Phone,Status\nAnn,Example,ann@example.test,555-555-0100,New\nAnn,Example,ann@example.test,555-555-0200,New\nJane,Other,jane@example.test,555-555-0100,New\n')
    response = preview(client, url, ['first_name', 'last_name', 'email', 'phone', 'status'], 'update')
    assert 'cannot remove an existing Do Not Call' in response.text
    assert 'multiple contacts' in response.text
    assert 'different name' in response.text
    commit(client, url)
    with sessions() as db:
        assert len(db.scalars(select(Lead)).all()) == 2
        assert db.scalar(select(Lead).where(Lead.email == 'ann@example.test')).status == 'Do Not Call'


def test_cookie_csrf_expiry_and_no_commit_without_preview(setup):
    client, sessions = setup
    url = stage(client, b'First,Last\nAnn,Example\n')
    assert client.post(url + '/commit', data={'confirm': 'yes'}).status_code == 403
    assert commit(client, url).status_code == 409
    other = TestClient(client.app)
    other.get('/imports')
    assert other.get(url).status_code == 404
    assert other.get(url + '/results.csv').status_code == 404
    with sessions.begin() as db:
        db.get(ContactImportBatch, url.rsplit('/', 1)[-1]).expires = datetime.utcnow() - timedelta(seconds=1)
    assert client.get(url).status_code == 410


def test_cancel_discards_without_contacts(setup):
    client, sessions = setup
    url = stage(client, b'First,Last\nAnn,Example\n')
    client.post(url + '/cancel', data={'csrf': client.cookies[COOKIE]})
    assert client.get(url).status_code == 404
    with sessions() as db:
        assert not db.scalars(select(Lead)).all()


def test_xlsx_tab_native_dates_and_leading_zero():
    book = Workbook()
    book.active.title = 'Ignore'
    sheet = book.create_sheet('Contacts')
    sheet.append(['First', 'Last', 'DOB', 'ZIP'])
    sheet.append(['Ann', 'Example', datetime(1961, 4, 15), '00501'])
    content = io.BytesIO()
    book.save(content)
    table = service.read_file('contacts.xlsx', content.getvalue(), 'Contacts')
    data, errors = service.validate_values(table['rows'][0]['values'], ['first_name', 'last_name', 'date_of_birth', 'zip_code'])
    assert not errors
    assert data['date_of_birth'] == date(1961, 4, 15)
    assert data['zip_code'] == '00501'
    with pytest.raises(service.ImportProblem):
        service.read_file('contacts.xlsx', content.getvalue(), 'Missing')


def test_xls_native_dates():
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('Contacts')
    for j, value in enumerate(['First', 'Last', 'DOB']):
        sheet.write(0, j, value)
    sheet.write(1, 0, 'Ann'); sheet.write(1, 1, 'Example')
    sheet.write(1, 2, datetime(1961, 4, 15), xlwt.easyxf(num_format_str='MM/DD/YYYY'))
    content = io.BytesIO(); book.save(content)
    table = service.read_file('contacts.xls', content.getvalue())
    assert table['rows'][0]['values'][2].startswith('1961-04-15')


@pytest.mark.parametrize('content', [b'A,A\n1,2', b'A,,C\n1,2,3', b'A,B\n1,2,3', b'', b'PKinvalid'])
def test_bad_csv_inputs(content):
    with pytest.raises(service.ImportProblem):
        service.read_file('contacts.csv', content)


def test_csv_encodings_quoted_fields_and_mapping_conflicts(setup):
    table = service.read_file('list.csv', 'First;Last;Notes\nAndr\u00e9;Example;"a; b"\n'.encode('cp1252'))
    assert table['rows'][0]['values'] == ['André', 'Example', 'a; b']
    assert service.suggested_mapping(['First', 'First Name', 'Last', 'MBI', 'PDP']) == ['first_name', '', 'last_name', 'medicare_mbi', 'interested_pdp']
    client, _ = setup
    url = stage(client, b'A,B\nAnn,Example\n')
    assert 'only one source column' in preview(client, url, ['first_name', 'first_name']).text
    assert 'both First Name and Last Name' in preview(client, url, ['first_name', '']).text


def test_google_sheet_private_reader_tab_selection_and_no_arbitrary_urls():
    class Reply:
        status_code = 200
        def __init__(self, data): self.data = data
        def json(self): return self.data
        def iter_content(self, size): yield json.dumps(self.data).encode()
        def close(self): pass
    class Google:
        def __init__(self): self.calls = []
        def get(self, url, **kwargs):
            self.calls.append((url, kwargs))
            if len(self.calls) == 1:
                return Reply({'sheets': [{'properties': {'sheetId': 0, 'title': 'Other'}}, {'properties': {'sheetId': 42, 'title': "Coach's Contacts"}}]})
            return Reply({'values': [['First', 'Last'], ['Ann', 'Example']]})
    google = Google()
    table = service.read_google_sheet('https://docs.google.com/spreadsheets/d/abc123/edit#gid=42', session=google)
    assert table['worksheet'] == "Coach's Contacts"
    assert "'Coach''s Contacts'" in unquote(google.calls[-1][0])
    assert google.calls[-1][1]['params']['valueRenderOption'] == 'FORMATTED_VALUE'
    assert google.calls[-1][1]['stream'] is True
    with pytest.raises(service.ImportProblem):
        service.read_google_sheet('https://example.com/spreadsheets/d/abc123/edit', session=google)
    with pytest.raises(service.ImportProblem):
        service.read_google_sheet('https://docs.google.com.evil.test/spreadsheets/d/abc123', session=google)


def test_google_error_missing_connection_and_source_route(setup, monkeypatch):
    monkeypatch.delenv('GOOGLE_SHEETS_CREDENTIALS_JSON', raising=False)
    monkeypatch.delenv('GOOGLE_APPLICATION_CREDENTIALS', raising=False)
    with pytest.raises(service.ImportProblem, match='not connected yet'):
        service.sheets_credentials()
    client, sessions = setup
    import app.imports_v35 as routes
    monkeypatch.setattr(routes, 'read_google_sheet', lambda link, tab: {'headers': ['First', 'Last'], 'rows': [{'row': 2, 'values': ['Ann', 'Example']}], 'worksheet': 'Contacts'})
    result = client.post('/imports/source', data={'csrf': client.cookies[COOKIE], 'source_type': 'sheets', 'sheet_url': 'https://docs.google.com/spreadsheets/d/abc/edit'})
    assert 'Google Sheets · Contacts' in result.text
    url = result.url.path
    preview(client, url, ['first_name', 'last_name'])
    commit(client, url)
    with sessions() as db:
        assert db.scalar(select(Lead)).first_name == 'Ann'


def test_validation_does_not_silently_change_bad_values():
    fields = ['first_name', 'last_name', 'email', 'state', 'date_of_birth', 'interested_pdp', 'status', 'call_attempts']
    _, errors = service.validate_values(['Ann', 'Example', 'no-email', 'Not a state', '04/05/61', 'sometimes', 'Unknown', '-2'], fields)
    assert len(errors) == 6


def test_row_limits_and_full_results(setup):
    with pytest.raises(service.ImportProblem):
        service.table_from_rows([['First', 'Last']] + [['Ann', 'Example']] * (service.MAX_ROWS + 1))
    client, _ = setup
    content = ('First,Last,Email\n' + '\n'.join(f'Ann,Example{i},a{i}@example.test' for i in range(205))).encode()
    url = stage(client, content)
    response = preview(client, url, ['first_name', 'last_name', 'email'])
    assert 'first 200 of 205' in response.text
    assert len(client.get(url + '/results.csv').text.splitlines()) == 206
    assert 'Import 205 contacts' in response.text


def test_templates_escape_source_data(setup):
    client, _ = setup
    url = stage(client, b'First,Last,Email\n<script>alert(1)</script>,Example,a@example.test\n')
    response = preview(client, url, ['first_name', 'last_name', 'email'])
    assert '<script>alert(1)</script>' not in response.text
    assert '&lt;script&gt;' in response.text


def test_reimport_same_source_skips_after_success(setup):
    client, sessions = setup
    content = b'First,Last,Email\nAnn,Example,ann@example.test\n'
    url = stage(client, content)
    preview(client, url, ['first_name', 'last_name', 'email'])
    commit(client, url)
    second = stage(client, content)
    preview(client, second, ['first_name', 'last_name', 'email'])
    commit(client, second)
    with sessions() as db:
        assert len(db.scalars(select(Lead)).all()) == 1


def test_name_only_reimport_needs_identifier(setup):
    client, sessions = setup
    content = b'First,Last\nAnn,Example\n'
    url = stage(client, content)
    preview(client, url, ['first_name', 'last_name'])
    commit(client, url)
    second = stage(client, content)
    assert 'no identifier to confirm a match' in preview(client, second, ['first_name', 'last_name']).text
    commit(client, second)
    with sessions() as db:
        assert len(db.scalars(select(Lead)).all()) == 1


def test_commit_failure_rolls_back_every_contact(setup, monkeypatch):
    client, sessions = setup
    url = stage(client, b'First,Last,Email\nAnn,Example,ann@example.test\nBob,Example,bob@example.test\n')
    preview(client, url, ['first_name', 'last_name', 'email'])
    from sqlalchemy import event
    def fail_second(mapper, connection, target):
        if target.first_name == 'Bob':
            raise ValueError('simulated database failure')
    event.listen(Lead, 'before_insert', fail_second)
    try:
        assert commit(client, url).status_code == 500
    finally:
        event.remove(Lead, 'before_insert', fail_second)
    with sessions() as db:
        assert db.scalars(select(Lead)).all() == []
        assert db.get(ContactImportBatch, url.rsplit('/', 1)[-1]).state == 'preview'
    assert commit(client, url).status_code == 200
    with sessions() as db:
        assert len(db.scalars(select(Lead)).all()) == 2


def test_competing_imports_require_fresh_preview(setup):
    client, sessions = setup
    content = b'First,Last,Email\nAnn,Example,ann@example.test\n'
    first, second = stage(client, content), stage(client, content)
    for url in [first, second]:
        preview(client, url, ['first_name', 'last_name', 'email'])
    commit(client, first)
    response = commit(client, second)
    assert 'Contacts changed since your preview' in response.text
    with sessions() as db:
        assert len(db.scalars(select(Lead)).all()) == 1


def test_old_preview_does_not_reopen_completed_import(setup):
    client, sessions = setup
    url = stage(client, b'First,Last,Email\nAnn,Example,ann@example.test\n')
    preview(client, url, ['first_name', 'last_name', 'email'])
    commit(client, url)
    assert 'Import complete' in preview(client, url, ['first_name', 'last_name', 'email']).text
    with sessions() as db:
        assert db.get(ContactImportBatch, url.rsplit('/', 1)[-1]).state == 'done'


def test_google_limits_and_permission_failures():
    class Reply:
        def __init__(self, data, status=200): self.data, self.status_code = data, status
        def json(self): return self.data
        def iter_content(self, size): yield json.dumps(self.data).encode()
        def close(self): pass
    class Google:
        def __init__(self, replies): self.replies = iter(replies)
        def get(self, *args, **kwargs): return next(self.replies)
    url = 'https://docs.google.com/spreadsheets/d/test/edit'
    with pytest.raises(service.ImportProblem, match='Cannot open'):
        service.read_google_sheet(url, session=Google([Reply({}, 403)]))
    metadata = Reply({'sheets': [{'properties': {'sheetId': 0, 'title': 'Contacts'}}]})
    with pytest.raises(service.ImportProblem, match='unique, nonblank'):
        service.read_google_sheet(url, session=Google([metadata, Reply({'values': [['First', 'First'], ['Ann', 'Example']]})]))
